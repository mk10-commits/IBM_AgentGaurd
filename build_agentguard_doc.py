from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r'C:\Users\Maryam\OneDrive\Desktop\IBM\AgentGuard_Hackathon_Blueprint_Implementation.docx'
BLUE = '0F62FE'
PALE = 'EAF2FF'
GRAY = 'F4F6F8'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(9.5); styles['Normal'].font.color.rgb = RGBColor(35, 42, 50)
styles['Normal'].paragraph_format.space_after = Pt(5)
for name, size in [('Title', 26), ('Heading 1', 17), ('Heading 2', 12)]:
    s=styles[name]; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(12 if name!='Title' else 0); s.paragraph_format.space_after=Pt(5)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def borders(table):
    tblPr=table._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); b.append(el)
    tblPr.append(b)
def set_cell(cell, text, bold=False, color=None, size=8.3):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2)
    r=p.add_run(text); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,h in enumerate(headers):
        set_cell(t.rows[0].cells[i],h,True,'FFFFFF',8.5); shade(t.rows[0].cells[i], '1F4E79')
    trPr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        if widths:
            for c,w in zip(cells,widths): c.width=Inches(w)
        for i,val in enumerate(row):
            set_cell(cells[i],str(val),False,None,8.1)
            if ridx%2==1: shade(cells[i],GRAY)
    borders(t); doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(text)
def numbered(text):
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(2); p.add_run(text)
def callout(label, text):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c=t.cell(0,0); shade(c,PALE); c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(3)
    r=p.add_run(label+'  '); r.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE)
    p.add_run(text); borders(t); doc.add_paragraph().paragraph_format.space_after=Pt(1)

# Cover / opening
p=doc.add_paragraph(style='Title'); p.add_run('AgentGuard Hackathon Blueprint')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(12)
r=p.add_run('Real-time action protection for autonomous AI agents'); r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(BLUE)
doc.add_paragraph('This blueprint describes a buildable AgentGuard prototype for the IBM Datathon. It covers the problem, priority features, system architecture, IBM technology integration, trade-offs, and a demo plan. The recommendation is to build a narrow, working action firewall for email, files, and payments, then prove that it blocks harmful actions while allowing normal work to continue.')
callout('Project idea', 'Before an AI agent can use a tool, AgentGuard checks what it wants to do, what information influenced the action, and how much harm it could cause. AgentGuard then allows, safely modifies, blocks, or escalates the action.')

doc.add_heading('1. Problem and project goal', level=1)
doc.add_paragraph('AI agents can read messages and documents, then act through tools such as email, databases, file systems, and payment services. A malicious instruction hidden in an email or an honest misunderstanding can lead the agent to send private information, delete data, or make an unauthorized payment. AgentGuard places an enforceable checkpoint between the agent and those tools.')
doc.add_paragraph('The prototype should demonstrate three things: the guard sees every action before it runs; it uses the user’s task and information sources to make a decision; and its decision keeps useful work moving where it can do so safely.')

doc.add_heading('2. Priority features and why they matter', level=1)
features=[
('1. Mandatory tool gateway','Every tool request passes through AgentGuard. The agent cannot call the underlying tool directly.','Without this boundary, an agent could bypass every other feature.'),
('2. Action normalization','Convert each request into a standard record: tool, action, target, data, destination, scope, and reversibility.','Makes the same policy usable across email, files, databases, and payment tools.'),
('3. Source and provenance tracking','Tag data as user/system input or as email, webpage, file, or tool output. Flag important action details that came only from untrusted sources. Show the source chain.','Helps identify indirect prompt injection and explains why an action looks suspicious.'),
('4. Scope, impact, and data sensitivity','Estimate the affected files/records/users, whether information is sensitive, whether the destination is external, and whether the action can be undone.','A single draft and a mass deletion should not receive the same treatment.'),
('5. Deterministic policy rules','Apply explicit rules for approved destinations, allowed paths, payment limits, secrets, and required approvals.','Fast, consistent, and easy to audit. Critical controls should not depend on an LLM judge.'),
('6. Risk score with reason codes','Combine impact, irreversibility, data sensitivity, destination trust, and provenance into a 0–1 score plus clear reasons.','A score helps prioritize, while reason codes make decisions understandable.'),
('7. Four decision outcomes','Allow safe actions; modify actions to reduce danger; block clearly disallowed actions; escalate uncertain or high-impact actions.','More useful than a simple allow/block filter.'),
('8. Safe modification','Replace permanent delete with recoverable move; prepare a payment instead of submitting it; redact secrets; reduce a bulk action to a smaller approved scope.','Preserves task progress while reducing possible harm. This should be a signature demo feature.'),
('9. Human approval queue','Show the proposed action, amount/scope, destination, data involved, provenance, risk reasons, and approve/deny controls.','Keeps people in control of irreversible actions without interrupting every routine task.'),
('10. Least-privilege tool access','Give the agent only the permissions needed for its task. For example, allow email drafting but reserve sending permission for the guard.','Limits damage if an attack succeeds or a model misunderstands.'),
('11. Audit trail','Record action, policy, verdict, reason codes, time, and human decision. Do not log raw secrets.','Supports investigation, accountability, and a strong live dashboard.'),
('12. Safety and usefulness evaluation','Measure attack success, false blocks, legitimate task completion, median and tail latency, and human escalation rates. Include benign lookalikes.','Shows that the system balances protection with practical usefulness.')]
table(['Feature','Improved implementation','Importance'],features,[1.4,3.45,2.1])

doc.add_heading('3. Decision policy', level=1)
doc.add_paragraph('Use hard rules for clear cases and risk thresholds for uncertain cases. Set stricter thresholds for actions that are difficult to reverse. An optional language-model judge can help with ambiguous intent checks, but it should not override database permissions or explicit safety rules.')
table(['Action category','Example','Suggested response'],[
('Low impact / reversible','Read a file; draft an email','Allow and log'),
('Moderate impact','Edit a file with a backup; update a small record','Allow or modify; snapshot first when possible'),
('High impact / external','Send email externally; bulk update; delete many files','Escalate or block based on policy'),
('Irreversible / financial','Submit payment; permanently delete data','Require explicit authorization; escalate when uncertain')],[1.55,2.65,2.75])
callout('Scoring note', 'A risk score is an estimate, not proof of safety. Keep clear rules for known dangerous actions, and evaluate the score against labeled examples before claiming it is calibrated.')

doc.add_heading('4. System architecture', level=1)
doc.add_paragraph('The agent proposes a tool call, but only the AgentGuard gateway can execute it. A policy service evaluates the request, and the selected verdict controls execution. The prototype should use mock tools so attacks and outcomes are repeatable and harmless.')
table(['Stage','What happens'],[
('1. User task','Store the original request as the trusted intent for this run.'),
('2. Agent','Plans and proposes a structured call such as send_email, delete_file, query_customer, or submit_payment.'),
('3. Normalize','Convert the call into a common action record with target, scope, data, destination, and reversibility.'),
('4. Provenance','Attach the sources for important values and context. Check whether an untrusted source introduced a recipient, URL, file, or instruction.'),
('5. Risk and policy','Run deterministic rules, sensitivity checks, and risk scoring. Use a small model judge only for unclear intent cases.'),
('6. Verdict','Allow, modify, block, or send to human approval.'),
('7. Tool execution','Execute only through the gateway using restricted credentials. Record the outcome.'),
('8. Dashboard and audit','Show action, verdict, reasons, source trace, latency, and human review result.')],[1.25,5.7])
callout('Critical design rule', 'The agent must not possess a second route to the underlying tools. If the guard can be bypassed, it cannot reliably enforce policy.')

doc.add_heading('5. IBM technologies in the design', level=1)
table(['IBM technology','Where it fits','Practical value'],[
('IBM LinuxONE','Host the gateway, policy engine, and audit services in a production-style deployment. IBM Secure Execution and confidential computing can isolate workloads and protect data while it is processed.','A secure runtime foundation for sensitive workloads. For the hackathon, only claim LinuxONE deployment if the prototype actually runs there.'),
('IBM Db2','Use as the protected customer/payment data store. Give AgentGuard a restricted database identity and use row/column access controls where available. Enable database auditing for access records.','Adds a data-layer control: the agent can only see authorized records and fields, even if it asks for too much.'),
('IBM watsonx.data','Optional governed data layer if the demo needs multiple data sources. Use access policies and governance metadata to constrain what data can be used.','Connects action-level decisions to governed data access.'),
('IBM DataStage','Optional pipeline for preparing benchmark data or consolidating activity logs for analysis.','Useful when there is a real data integration need; not required for the core prototype.'),
('IBM watsonx.ai / Granite Guardian','Optional second-opinion judge for ambiguous intent checks, such as whether a proposed action serves the original task.','Can improve context-based review, but should never replace deterministic rules or authorization checks.')],[1.7,3.2,2.05])
doc.add_paragraph('Recommended datathon scope: implement the gateway and policy logic first using mock tools; use Db2 if the team can access it; and describe LinuxONE as the intended secure deployment environment unless the demo is actually hosted there. DataStage and watsonx.data are optional extensions, not prerequisites.')

doc.add_heading('6. Benefits, disadvantages, and mitigations', level=1)
table(['Benefit or limitation','What it means','How to address it'],[
('Benefit: stops actions before side effects','The guard can block a send, delete, or payment before a tool executes.','Demonstrate the blocked tool call and confirm that the mock system did not change.'),
('Benefit: safer continuation','The guard can turn a dangerous action into a recoverable or approval-based one.','Show the modified action and its final effect.'),
('Benefit: explainable decisions','Provenance and reason codes give a reviewer evidence for the verdict.','Use concise, visible reasons in the dashboard.'),
('Limitation: false blocks','Legitimate tasks may resemble attacks.','Include benign lookalikes and report false-block rate.'),
('Limitation: latency','Every tool call adds a check. Model-based checks can be slower.','Keep deterministic checks inline; call the model judge only for ambiguous cases; report measured latency.'),
('Limitation: incomplete provenance','A prototype can track tagged sources but cannot perfectly know every influence on a model.','State the scope clearly and show source tags from the wrapper.'),
('Limitation: bypass risk','A direct tool connection can avoid the gateway.','Use one execution path and restricted tool credentials.'),
('Limitation: judge errors','An LLM judge can miss an attack or overreact.','Use it as an advisory signal; keep authorization and high-risk rules deterministic.'),
('Limitation: IBM environment access','LinuxONE and enterprise services may not be available to the team.','Make the prototype portable with mock services and accurately distinguish tested components from planned integrations.')],[1.45,2.65,2.85])

doc.add_heading('7. Hackathon demo and evaluation', level=1)
doc.add_paragraph('Use a three-minute scenario called “Summarize my inbox and draft replies.” The fake inbox, file system, and banking API make the demo safe and deterministic.')
for x in [
'Normal email: the agent summarizes a message and drafts a reply. AgentGuard allows these low-risk actions.',
'Prompt injection: an email tells the agent to send the contact list to an unknown address. AgentGuard highlights that the address came from the email and blocks the external send.',
'Bulk deletion: the agent tries to permanently delete many files. AgentGuard converts the request into a recoverable move or asks for approval.',
'Payment: the agent proposes a payment above the configured threshold. AgentGuard shows the amount and reasons in the human approval queue.',
'Scoreboard: compare no guard, simple keyword rules, and full AgentGuard using the same scenarios.'
]: numbered(x)
doc.add_heading('Metrics to show', level=2)
table(['Metric','What it tells judges'],[
('Attack success rate','How often the attack still causes the harmful tool action.'),
('False-block rate','How often the guard blocks a legitimate task.'),
('Task completion retention','Whether the agent still completes normal work.'),
('Added latency','How much delay the guard adds per tool call.'),
('Human escalation rate','How frequently the guard needs a person, and how often the person agrees.')],[1.9,5.05])
doc.add_paragraph('Run multiple trials for model-driven cases and label results honestly. Do not invent performance numbers. If the dataset or benchmark license is unclear, use your own synthetic cases until you verify the terms.')

doc.add_heading('8. Build priorities', level=1)
table(['Priority','Build item','Demo evidence'],[
('Must build','Mock tools; mandatory gateway; action normalization; allow/block/modify/escalate; basic provenance tags; simple rules; event log.','Show all four verdicts and prove the mock tool only changes after an allow or safe modification.'),
('Next','Risk score and reasons; approval queue; redaction; benign lookalikes; benchmark comparison dashboard.','Show the decision logic and quantify protection versus usefulness.'),
('If time permits','Db2 integration; Granite Guardian for ambiguous cases; LinuxONE deployment; DataStage or watsonx.data integration.','Show actual integration evidence, or clearly label it as a future deployment option.')],[1.1,3.2,2.65])

doc.add_heading('9. Proposed pitch', level=1)
callout('Pitch', 'AgentGuard checks what an AI agent is about to do, where the instruction came from, and how much harm it could cause. It allows safe actions, makes risky actions safer, blocks dangerous ones, and asks a person to approve high-impact decisions.')

doc.add_heading('10. Recommended implementation stack', level=1)
doc.add_paragraph('Choose a stack that lets the team finish a reliable vertical slice quickly. For the hackathon, a simple Python service and a small dashboard are more valuable than a broad enterprise deployment with unfinished integrations.')
table(['Layer','Recommended choice','How to use it'],[
('Language and API','Python 3.11+ with FastAPI','Build the guard as a small HTTP service. FastAPI provides a straightforward way to expose endpoints and interactive API documentation.'),
('Request validation','Pydantic models','Validate tool-call payloads and verdicts before policy evaluation or execution. Reject malformed or unknown action types.'),
('Agent connector','A thin adapter around plain tool calling; optionally LangGraph if already familiar','Keep the guard independent of the agent framework. The adapter converts the framework’s tool request to the AgentGuard schema.'),
('Policy and risk engine','Python modules with explicit rule functions','Keep rules deterministic, unit-sized, and easy to explain. Start with allowlists, deny rules, path scope, external recipients, data sensitivity, reversibility, and transaction limits.'),
('Demo data store','SQLite for the fastest local demo; Db2 as the IBM integration option','Store mock inbox items, files, payment requests, approval decisions, and audit events. Never use real customer or payment data.'),
('Dashboard','Streamlit for fastest build; React only if the team already has frontend capacity','Show an event stream, risk reasons, source trail, and approval queue. Streamlit reduces frontend setup for a short event.'),
('Optional model judge','watsonx.ai API with a supported Granite Guardian model','Call only when policy rules cannot confidently decide. Set a timeout and fail safely for high-impact actions.'),
('Packaging','Docker Compose','Run the API, dashboard, and optional database with one repeatable command. Keep a mock-only mode for a reliable live demo.'),
('IBM production runtime','LinuxONE with IBM Secure Execution for Linux, when accessible','Deploy the containerized service inside a protected Linux guest. Secure Execution protects guest memory/state from the hosting environment; persistent storage and application authorization still need their own controls.')],[1.45,2.15,3.35])
doc.add_paragraph('Why this stack: FastAPI and Pydantic give a clear typed boundary for every proposed action; a local database makes the demo repeatable; and Streamlit keeps attention on the security decisions rather than frontend plumbing. IBM’s watsonx.ai REST API supports programmatic foundation-model inference. IBM Secure Execution for Linux protects secure KVM guest workloads from host inspection or modification, but it is a deployment layer—not a replacement for AgentGuard policies, database access control, encrypted storage, or secret management.')

doc.add_heading('11. Detailed implementation steps', level=1)
steps=[
('Step 1 — Freeze the demo scope','Pick three tool families: email, files, and payment. Define a handful of safe and unsafe operations for each. Use mock data and a fake bank endpoint. Keep the first version framework-neutral.'),
('Step 2 — Define the action contract','Create a Pydantic ToolAction model with request_id, run_id, actor_id, tool_name, action, arguments, target, scope_count, reversibility, data_labels, destination, and provenance_refs. Create a GuardDecision model with verdict, risk_score, reason_codes, explanation, modified_action, approval_required, and policy_version.'),
('Step 3 — Build the mock tools first','Implement functions for read_email, draft_email, send_email, list_files, move_to_trash, delete_file, and submit_payment. Each function should make a visible state change so the dashboard can prove whether an action executed.'),
('Step 4 — Put the gateway in front','Build POST /v1/actions/evaluate-and-execute. It receives the proposed action, validates it, evaluates policy, then either executes the approved tool, returns a modified action for execution, records a block, or places the action in the approval queue. Do not expose direct tool credentials to the agent.'),
('Step 5 — Add provenance tags','When the mock inbox or file system returns content, attach a source_id, source_type, trust_level, and parent_source_ids. Carry those references with derived values, such as an email recipient copied from a message. In the prototype, describe this accurately as wrapper-based provenance tracking, not perfect causal understanding of the model.'),
('Step 6 — Add deterministic controls','Implement explicit checks before any model judge: tool and action allowlist; path boundary; external destination allowlist; secret/PII scan; bulk operation limit; payment cap; and required-user-authorization check for irreversible actions. Return structured reason codes instead of free-form text alone.'),
('Step 7 — Implement risk and decision policy','Normalize features to a 0–1 scale. For a prototype, try R = 0.25I + 0.20U + 0.20S + 0.20P + 0.15D, where I is impact, U is irreversibility, S is sensitivity, P is untrusted provenance involvement, and D is destination risk. Treat this only as an initial heuristic. Let hard rules override the score. Use action-specific outcomes: low-risk read/draft can pass; edits should snapshot; destructive or external actions may be modified or escalated; clearly disallowed actions are blocked.'),
('Step 8 — Make modification executable','Do not just tell the agent an action is safer. Produce a concrete modified action, show the before/after in the UI, and pass only the modified action to the mock tool. Examples: delete_file becomes move_to_trash; send_email with a secret becomes redact_then_draft; payment becomes create_payment_draft.'),
('Step 9 — Add human approval','Store escalated actions with pending status. Provide approve and deny controls. On approval, execute the exact reviewed payload, not a newly generated one. If the payload changes, require a new review. Include approver and decision time in the audit event.'),
('Step 10 — Add audit events and dashboard','Record request ID, action summary, verdict, reason codes, risk score, policy version, elapsed milliseconds, execution result, and approval result. Mask secrets and personal values in logs. Dashboard cards should show allowed, modified, blocked, awaiting approval, attack success, false blocks, and latency.'),
('Step 11 — Add the optional IBM model check','Only after the deterministic path works, connect an IBM Granite Guardian model through watsonx.ai for ambiguous intent cases. Send the minimum required context, set a strict timeout, record model/version metadata, and route uncertain results to a human for high-impact actions. If credentials or network access are unavailable, keep this as a clearly labeled integration point.'),
('Step 12 — Validate with a fixed scenario set','Create paired cases: injected external send versus normal reply; mass delete versus cleanup of one temporary file; over-limit payment versus permitted payment; secret leak versus safe redacted message. Compare no guard, rules only, and full guard. Repeat stochastic agent runs and report measured results, including false blocks and latency.'),
('Step 13 — Prepare a reliable presentation','Seed the mock environment before the demo, rehearse the exact three-minute flow, show one live attack and one safe modification, then finish on the metrics panel. Keep a screen recording as a fallback. Clearly separate implemented integrations from planned IBM deployment options.')]
for title,body in steps:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(title+'. '); r.bold=True; p.add_run(body)

doc.add_heading('12. API and data design', level=1)
doc.add_paragraph('Keep the external request format small and stable so a LangGraph, CrewAI, or custom agent adapter can submit the same action shape.')
table(['Field','Example / meaning'],[
('tool_name / action','email / send_email; files / delete_file; payments / submit_payment'),
('target and scope','Recipient or resource path; number of files/records affected'),
('arguments','Tool-specific structured arguments; redact secrets before storing'),
('data_labels','Public, internal, personal, confidential, secret'),
('provenance_refs','IDs linking recipient, URL, or content to user input or external source'),
('reversibility','reversible, recoverable, partially_reversible, irreversible'),
('decision','allow, modify, block, escalate; with score and reason codes'),
('execution status','not_run, completed, failed, awaiting_approval, denied')],[1.8,5.15])
table(['Endpoint','Purpose'],[
('POST /v1/actions/evaluate-and-execute','Validate, evaluate, and execute or route the proposed action.'),
('GET /v1/events?run_id=...','Retrieve the run’s event stream for the dashboard.'),
('GET /v1/approvals','List pending human reviews.'),
('POST /v1/approvals/{id}/decision','Approve or deny the reviewed payload.'),
('POST /v1/demo/reset','Reset the mock environment before a rehearsal.')],[2.55,4.4])

doc.add_heading('13. Suggested repository structure', level=1)
for line in [
'agentguard/ — application package',
'  api/ — FastAPI routes and request/response models',
'  schemas/ — Pydantic ToolAction, Provenance, and GuardDecision models',
'  normalizer/ — framework adapters and action classification',
'  provenance/ — source tagging and lineage references',
'  policies/ — deterministic rules, risk features, thresholds, reason codes',
'  tools/ — mock email, file, Db2, and payment tools',
'  approvals/ — pending review storage and decision handling',
'  audit/ — event writer and redaction-safe log formatting',
'  integrations/ — optional watsonx.ai and Db2 clients',
'dashboard/ — Streamlit event stream and approval queue',
'data/ — synthetic demo fixtures only',
'scenarios/ — benign and adversarial scripted demo cases',
'infra/ — Docker Compose and deployment configuration',
'README.md — setup, run, demo script, limitations, and measured results']:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0)
    r=p.add_run(line); r.font.name='Consolas'; r.font.size=Pt(8.5)

doc.add_heading('14. Suggested 24-hour execution plan', level=1)
table(['Time window','Deliverable','Completion check'],[
('Hours 0–2','Select scenarios, agree on action schema, create repository and mock data.','A clear story and scripted cases work on paper.'),
('Hours 2–6','Mock tools and gateway; make tool execution possible only through the guard.','One allow and one block can be shown end to end.'),
('Hours 6–10','Normalization, source labels, deterministic rules, reason codes.','Injected recipient and destructive action are caught with visible reasons.'),
('Hours 10–14','Modify and approval paths; event store.','Trash conversion and payment approval operate on mock state.'),
('Hours 14–17','Dashboard with event stream and approval queue.','Judges can see decisions and side effects without reading terminal output.'),
('Hours 17–20','Paired benign/adversarial evaluation and baseline comparison.','Metrics are computed from recorded outcomes, not guessed.'),
('Hours 20–22','Optional Db2 or watsonx.ai connection if available.','Core demo remains functional if integration is switched off.'),
('Hours 22–24','Rehearsal, slides, latency capture, and fallback recording.','The demo resets cleanly and fits the event time limit.')],[1.15,3.0,2.8])

doc.add_heading('References', level=1)
refs=[
('IBM Confidential Computing Platform','https://www.ibm.com/products/confidential-computing-platform'),
('IBM Secure Execution for Linux overview','https://www.ibm.com/docs/en/linux-on-systems?topic=execution-introduction'),
('IBM Db2 row and column access control','https://www.ibm.com/docs/en/db2/12.1.0?topic=security-row-column-access-control-rcac'),
('IBM Db2 audit facility','https://www.ibm.com/docs/en/db2/11.5.x?topic=activities-introduction-db2-audit-facility'),
('IBM watsonx.data access management and governance','https://www.ibm.com/docs/en/watsonxdata/saas?topic=components-access-management-governance'),
('IBM DataStage','https://www.ibm.com/products/datastage'),
('IBM watsonx.ai foundation models','https://www.ibm.com/products/watsonx-ai/foundation-models'),
('IBM watsonx.ai REST API','https://www.ibm.com/docs/en/watsonx/saas?topic=resources-rest-api'),
('FastAPI first steps','https://fastapi.tiangolo.com/tutorial/first-steps/'),
('CaMeL research on prompt-injection defenses','https://arxiv.org/abs/2503.18813'),
('AgentDojo benchmark paper','https://proceedings.neurips.cc/paper_files/paper/2024/file/97091a5177d8dc64b1da8bf3e1f6fb54-Paper-Datasets_and_Benchmarks_Track.pdf')]
for title,url in refs:
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(1)
    r=p.add_run(title+' — '+url); r.font.size=Pt(8); r.font.color.rgb=RGBColor(31,78,121)

# Footer with page field
for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=footer.add_run('AgentGuard | IBM Datathon  •  '); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(100,110,120)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

doc.core_properties.title='AgentGuard Hackathon Blueprint'
doc.core_properties.subject='Runtime action firewall for autonomous AI agents'
doc.core_properties.author=''
doc.save(OUT)
print(OUT)


